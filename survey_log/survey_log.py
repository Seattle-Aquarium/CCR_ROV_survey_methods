"""The ROV survey flight log: tally it, or re-export it from the tracking spreadsheet.

    python survey_log/survey_log.py
        Tally ROV_survey_flights.csv: transects and survey days, by year,
        program and region. Standard library only.

    python survey_log/survey_log.py export --xlsx path/to/3_ROV_flight_log.xlsx
        Rewrite ROV_survey_flights.csv from the spreadsheet's ROV_survey_flights
        tab (needs openpyxl). The People, Notes and issues columns are left out.

Counting rules (the spreadsheet's "Counted as survey transect" and
"Transects counted" columns carry them row by row):

* A transect is a row with a transect number. Rows marked NA (no transect
  flown) or error (aborted) do not count, and neither do test, outreach and
  partner-demonstration flights or short aborted dives. A transect flown in
  two parts counts once.
* On diver-ROV comparison days the ROV's passes along each 100 m tape were
  processed as three 30 m transects, matching the divers' Reef Check
  transects, so those days count six transects however many passes the
  field log records.
* A survey day is a calendar date on which at least one transect counted.
"""
import argparse
import csv
import re
from collections import Counter, defaultdict
from datetime import date, datetime, time
from pathlib import Path

HERE = Path(__file__).resolve().parent
CSV_PATH = HERE / "ROV_survey_flights.csv"
FIELDS = ["date", "program", "region", "port_site", "location", "transect", "vehicle",
          "start", "end", "pauses", "counted", "transects_counted", "note"]

# Days logged under "Elliott Bay" that belong to the HSIL program, not the Port's.
HSIL_ELLIOTT_BAY_DAYS = {"2025-08-28", "2025-08-29"}
PORT_LOCATIONS = ("elliott bay marina", "pocket beach", "sirens of spring")
REGIONS = [  # first match wins, on the lower-cased location
    ("jefferson head and centennial", "Central Puget Sound; Elliott Bay"),
    ("jefferson", "Central Puget Sound"), ("kingston", "Central Puget Sound"),
    ("shaw", "San Juan Islands"), ("satellite", "San Juan Islands"),
    ("deadman", "San Juan Islands"), ("friday harbor", "San Juan Islands"),
    ("anacortes", "Fidalgo Island"),
    ("neah bay", "Strait of Juan de Fuca"), ("tatoosh", "Strait of Juan de Fuca"),
    ("koitlah", "Strait of Juan de Fuca"), ("mushroom", "Strait of Juan de Fuca"),
    ("site 5", "Strait of Juan de Fuca"),
    ("cathlamet", "Other"),
]


def program_of(site, location, day):
    s = str(site or "").strip()
    if re.fullmatch(r"\d+", s):
        return "Port of Seattle"
    if s in ("HSIL", "Olympic Coast"):
        return s
    if s == "Elliott Bay":
        return "HSIL" if day in HSIL_ELLIOTT_BAY_DAYS else "Port of Seattle"
    loc = location.lower()
    if "neah bay" in loc:
        return "Olympic Coast"
    if any(p in loc for p in PORT_LOCATIONS):
        return "Port of Seattle"
    return "Other"


def region_of(location):
    loc = location.lower()
    return next((r for key, r in REGIONS if key in loc), "Elliott Bay")


def export(xlsx):
    import openpyxl

    ws = openpyxl.load_workbook(xlsx, data_only=True)["ROV_survey_flights"]
    anchor = {}
    for rng in ws.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                anchor[(r, c)] = (rng.min_row, rng.min_col)

    def val(r, c):
        return ws.cell(*anchor.get((r, c), (r, c))).value

    def text(v):
        if v is None:
            return ""
        if isinstance(v, datetime):
            return v.date().isoformat()
        if isinstance(v, (date, time)):
            return v.isoformat()
        return str(v).strip()

    header = [str(ws.cell(1, c).value or "").strip() for c in range(1, ws.max_column + 1)]
    col = {name: i + 1 for i, name in enumerate(header)}
    rows = []
    for r in range(2, ws.max_row + 1):
        day = text(val(r, col["Date"]))
        if not day:
            continue
        site, loc = val(r, col["Site"]), text(val(r, col["Location"]))
        verdict = text(val(r, col["Counted as survey transect"]))
        counted, _, note = verdict.partition(" - ")
        rows.append({
            "date": day,
            "program": program_of(site, loc, day),
            "region": region_of(loc),
            "port_site": text(site) if re.fullmatch(r"\d+", text(site)) else "",
            "location": loc,
            "transect": text(val(r, col["Transect #"])),
            "vehicle": text(val(r, col["ROV"])),
            "start": text(val(r, col["Start Time"])),
            "end": text(val(r, col["End Time"])),
            "pauses": text(val(r, col["Pauses"])),
            "counted": counted,
            "transects_counted": int(val(r, col["Transects counted"]) or 0),
            "note": note,
        })
    with open(CSV_PATH, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(rows)
    print(f"wrote {len(rows)} rows to {CSV_PATH}")


def tally(path=CSV_PATH):
    with open(path, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    counted = [r for r in rows if int(r["transects_counted"]) > 0]
    total = sum(int(r["transects_counted"]) for r in counted)
    days = {r["date"] for r in counted}
    print(f"{len(rows)} logged rows; {total} survey transects on {len(days)} survey days, "
          f"{min(days)} to {max(days)}")

    def group(key):
        t, d = Counter(), defaultdict(set)
        for r in counted:
            k = key(r)
            t[k] += int(r["transects_counted"])
            d[k].add(r["date"])
        return t, d

    for title, key in (("year", lambda r: r["date"][:4]), ("program", lambda r: r["program"]),
                       ("region", lambda r: r["region"])):
        t, d = group(key)
        print(f"\nby {title}:")
        for k in sorted(t):
            print(f"  {k:<34} {len(d[k]):3d} days  {t[k]:4d} transects")
    reasons = Counter(r["note"] for r in rows if r["counted"] == "no")
    print("\nrows not counted:")
    for k, v in reasons.most_common():
        print(f"  {v:3d}  {k}")
    missing = sum(1 for r in counted if not (r["start"] and r["end"]))
    print(f"\ncounted rows without logged start/end times: {missing}")


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    sub = ap.add_subparsers(dest="cmd")
    ex = sub.add_parser("export", help="rewrite the CSV from the tracking spreadsheet")
    ex.add_argument("--xlsx", type=Path, required=True)
    args = ap.parse_args()
    if args.cmd == "export":
        export(args.xlsx)
    tally()


if __name__ == "__main__":
    main()
